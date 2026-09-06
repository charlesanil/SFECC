import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


# Project paths
PROJECT_ROOT = Path(__file__).resolve().parent
EXCEL_FILE = PROJECT_ROOT / "data" / "SFECC_Environmental_Case_Tracker_v1.xlsx"
JSON_FILE = PROJECT_ROOT / "data" / "ngt-cases.json"


# Only these six fields are published on the SFECC website
PUBLIC_FIELDS = {
    "Case No": "caseNo",
    "Case Title": "caseTitle",
    "Bench": "bench",
    "District": "district",
    "Current Status": "status",
    "Registt Date": "registrationDate",
}


def clean_text(value):
    """Clean normal text fields."""
    if value is None:
        return ""

    return " ".join(str(value).split())


def clean_date(value):
    """Convert dates to YYYY-MM-DD."""
    if value is None:
        return ""

    # Real Excel date/datetime
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")

    value = str(value).strip()

    # Dates stored as text in Excel
    for date_format in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            date = datetime.strptime(value, date_format)
            return date.strftime("%Y-%m-%d")
        except ValueError:
            pass

    # If the value doesn't look like a date, preserve it
    return value


def main():
    print("Reading Excel file...")
    print(f"Source: {EXCEL_FILE}")

    if not EXCEL_FILE.exists():
        print("ERROR: Excel file not found.")
        return

    workbook = load_workbook(EXCEL_FILE, data_only=True)

    if "Case Register" not in workbook.sheetnames:
        print("ERROR: 'Case Register' sheet not found.")
        return

    sheet = workbook["Case Register"]

    # Read Excel headers
    headers = [cell.value for cell in sheet[1]]

    # Check required columns
    missing = [field for field in PUBLIC_FIELDS if field not in headers]

    if missing:
        print("ERROR: Missing columns in Excel:")
        for field in missing:
            print(f"  - {field}")
        return

    header_index = {
        header: index
        for index, header in enumerate(headers)
    }

    cases = []

    for row in sheet.iter_rows(min_row=2, values_only=True):

        # Skip completely empty rows
        if not any(value is not None for value in row):
            continue

        case = {}

        for excel_field, json_field in PUBLIC_FIELDS.items():
            value = row[header_index[excel_field]]

            if excel_field == "Registt Date":
                case[json_field] = clean_date(value)
            else:
                case[json_field] = clean_text(value)

        cases.append(case)

    # Write public JSON
    with open(JSON_FILE, "w", encoding="utf-8") as file:
        json.dump(
            cases,
            file,
            indent=2,
            ensure_ascii=False
        )

    print()
    print(f"SUCCESS: {len(cases)} cases exported.")
    print(f"Public JSON: {JSON_FILE}")


if __name__ == "__main__":
    main()